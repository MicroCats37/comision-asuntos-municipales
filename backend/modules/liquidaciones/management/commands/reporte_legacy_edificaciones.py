"""
Management command to generate a detailed CSV report comparing legacy Excel values
against `cotizar_legacy_proceso` calculations — WITHOUT inserting anything.

Usage:
    python manage.py reporte_legacy_edificaciones --settings=config.settings.development
    python manage.py reporte_legacy_edificaciones --data-path "C:\\...\\data.csv" --settings=config.settings.development
    python manage.py reporte_legacy_edificaciones --output "reporte_detalle.csv" --settings=config.settings.development

CSV columns (one row per liquidacion, exact order):
    FILA, NRO, EXPEDIENTE, FECHA_REGISTRO, ESPECIALIDAD, RAZON_SOCIAL,
    VALOR_OBRA, PORCENTAJE_EXCEL, PORCENTAJE_CALCULADO,
    SUBTOTAL_EXCEL, SUBTOTAL_CALCULADO, TOTAL_EXCEL, TOTAL_CALCULADO,
    DESCRIPCION_LEGACY, ESTADO

ESTADO rules:
    RECHAZADO  — NRO == 0  OR  DPTOPRDI (distrito) empty
    ERROR      — exception raised during payload building / cotizar call
    AVISO      — uploaded (not rejected/error) but anomalies present
    OK         — uploaded with no anomalies
"""

import csv
import io
import logging
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.core.exceptions import ObjectDoesNotExist
from django.core.management.base import BaseCommand, CommandError
from injector import Injector

from modules.entidades.domain.models.municipalidad import Municipalidad
from modules.entidades.domain.models.ubigeo import UbigeoDistrito
from modules.liquidaciones.di import LiquidacionesModule
from modules.liquidaciones.domain.services.orchestrators.liquidacion_legacy.liquidacion_edificaciones_legacy_orchestrator import (
    LiquidacionEdificacionesLegacyOrchestrator,
)
from modules.liquidaciones.presentation.schemas.liquidacion_general.general_schemas import (
    EntidadInlineSchema,
    ProyectoCotizarSchema,
)
from modules.liquidaciones.presentation.schemas.liquidacion_legacy.liquidacion_edificaciones_legacy_schemas import (
    LiquidacionEdificacionesLegacyIn,
    LiquidacionGeneralLegacyIn,
)
from modules.liquidaciones.presentation.schemas.liquidacion_tipo.porcentaje_schemas import (
    LiquidacionPorcentajeObraDatosIn,
    LiquidacionPorcentajeObraIn,
)
from modules.usuarios.di import UsuariosModule

logger = logging.getLogger(__name__)

# Default path to legacy data file
DEFAULT_DATA_PATH = Path(__file__).resolve().parent.parent.parent.parent.parent / "data-old" / "data.csv"

# openpyxl column indices — 0-indexed (list(row) from iter_rows(values_only=True))
# Header (row 1): ID=0 NRO=1 RUC=2 NOMBRE=3 PROYECTO=4 DPTOPRDI=5 DIRECCION=6
# VALOROBRA=7 PORCENTAJE=8 SUBTOTAL=9 IGV=10 TOTAL=11 FECHA=12 DNI=13
# RAZONSOCIAL=14 NROFACTURA=15 CODPAGO=16 NROEXPDTE=17 NROREV=18 ESPECIALIDAD=19
COL_RUC = 2
COL_NRO = 1
COL_NOMBRE = 3
COL_PROYECTO = 4
COL_DPTOPRDI = 5
COL_DIRECCION = 6
COL_VALOROBRA = 7
COL_PORCENTAJE = 8
COL_SUBTOTAL = 9
COL_IGV = 10
COL_TOTAL = 11
COL_FECHA = 12
COL_DNI = 13
COL_RAZONSOCIAL = 14
COL_CODPAGO = 16
COL_NROEXPDTE = 17
COL_NROREV = 18
COL_ESPECIALIDAD = 19

MAX_REVISIONES = 5


# ── Top-level helpers (also used by import_legacy_edificaciones) ──────────────

def _parse_fecha(raw) -> date:
    """Parse FECHA value (datetime object or string) -> date."""
    if not raw:
        raise ValueError("FECHA is empty")
    if isinstance(raw, datetime):
        return raw.date()
    raw_str = str(raw).strip()
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw_str, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Cannot parse fecha: {raw_str!r}")


def _resolve_tipo_documento(dni: str) -> tuple[str, str]:
    """
    Return (tipo_documento, numero_documento) based on DNI content.

    Rules:
        - empty/blank -> SIN_DOCUMENTO, 00000000
        - 8 chars (all digits) -> DNI, the value
        - otherwise -> RUC, the value
    """
    dni = str(dni).strip() if dni else ""
    if not dni:
        return "SIN_DOCUMENTO", "00000000"
    if len(dni) == 8 and dni.isdigit():
        return "DNI", dni
    return "RUC", dni


_DISTRITO_SINONIMOS = [
    (r"^CERCADO\s+DE\s+LIMA$", ("LIMA", None)),
    (r"^LIMA\s+CERCADO\s+Y\s+PROVINCIAL\s+LIMA$", ("LIMA", None)),
    (r"^ATE\s+VITARTE$", ("ATE", None)),
    (r"^LURIGANCHO\s*[-–]\s*CHOSICA$", ("LURIGANCHO", None)),
    (r"^BARRANCA\s*[-–]\s*NORTE$", ("BARRANCA", None)),
    (r"^SAN\s+ANTONIO\s+DE\s+HUAROCHIRI$", ("SAN ANTONIO", "HUAROCHIRI")),
]


def _normalizar_distrito(nombre: str) -> tuple[str, str | None]:
    """Normaliza un nombre histórico de distrito a (nombre canónico, provincia o None)."""
    if not nombre:
        return "", None
    nombre = nombre.strip().upper()
    for pattern, (canon, provincia) in _DISTRITO_SINONIMOS:
        if re.match(pattern, nombre):
            return canon, provincia
    return nombre, None


def _build_orchestrator() -> LiquidacionEdificacionesLegacyOrchestrator:
    """Build a fully-injected LiquidacionEdificacionesLegacyOrchestrator."""
    injector = Injector([LiquidacionesModule(), UsuariosModule()])
    return injector.get(LiquidacionEdificacionesLegacyOrchestrator)


# ── _read_excel — replicated from import_legacy_edificaciones.Command ──────────

def _read_excel(path: Path) -> list:
    """
    Read an Excel .xlsx file with openpyxl.

    The source file has a .csv extension but is actually an Excel .xlsx file.
    openpyxl rejects the extension, so we read as binary and load via BytesIO.
    """
    with path.open("rb") as f:
        data = f.read()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


# ── Main Command ───────────────────────────────────────────────────────────────

class Command(BaseCommand):
    help = "Generate a detailed CSV report comparing legacy Excel vs calculated values"

    def add_arguments(self, parser):
        parser.add_argument(
            "--data-path",
            type=str,
            default=None,
            help=f"Ruta al archivo Excel (default: {DEFAULT_DATA_PATH})",
        )
        parser.add_argument(
            "--output",
            type=str,
            default=None,
            help="Ruta del CSV de salida (default: reporte_legacy_detallado.csv junto al data-path)",
        )

    def handle(self, *args, **options):
        data_path = Path(options["data_path"]) if options["data_path"] else DEFAULT_DATA_PATH
        if not data_path.exists():
            raise CommandError(f"Data file not found: {data_path}")

        output_path = Path(options["output"]) if options["output"] else data_path.parent / "reporte_legacy_detallado.csv"

        self._log("\n[reporte_legacy_edificaciones] Starting...")
        self._log(f"  Source: {data_path}")
        self._log(f"  Output: {output_path}")

        # Resolve orchestrator
        self._log("  Resolving orchestrator...")
        orchestrator = _build_orchestrator()
        self._log("  Orchestrator ready.")

        # Read Excel
        rows = _read_excel(data_path)
        total_rows = len(rows)
        self._log(f"  Total rows: {total_rows}")

        # Open output CSV
        out_file = output_path.open("w", newline="", encoding="utf-8")
        writer = csv.writer(out_file)
        writer.writerow([
            "FILA", "NRO", "EXPEDIENTE", "FECHA_REGISTRO", "ESPECIALIDAD",
            "RAZON_SOCIAL", "VALOR_OBRA",
            "PORCENTAJE_EXCEL", "PORCENTAJE_CALCULADO",
            "SUBTOTAL_EXCEL", "SUBTOTAL_CALCULADO",
            "TOTAL_EXCEL", "TOTAL_CALCULADO",
            "DESCRIPCION_LEGACY", "ESTADO",
        ])

        # Counters
        counts = {"OK": 0, "AVISO": 0, "RECHAZADO": 0, "ERROR": 0}

        for idx, row in enumerate(rows, start=2):
            fila = idx
            estado = "ERROR"  # default until proven otherwise
            anomalies: list[str] = []

            # Column access helper — row is bound explicitly per iteration
            def col(r: list, n: int):
                return r[n] if len(r) > n else None

            nro_raw = col(row, COL_NRO)
            try:
                nro_val = int(float(str(nro_raw).strip())) if nro_raw is not None and str(nro_raw).strip() != "" else None
            except (ValueError, TypeError):
                nro_val = None

            dptoprdri = col(row, COL_DPTOPRDI)

            # ── RECHAZADO ───────────────────────────────────────────────────────
            if nro_val == 0 or (dptoprdri is None or str(dptoprdri).strip() == ""):
                estado = "RECHAZADO"
                if nro_val == 0:
                    anomalies.append("NRO == 0")
                if dptoprdri is None or str(dptoprdri).strip() == "":
                    anomalies.append("Distrito vacío (DPTOPRDI)")

                writer.writerow([
                    fila,
                    nro_val if nro_val is not None else "",
                    "",  # EXPEDIENTE
                    "",  # FECHA_REGISTRO
                    col(row, COL_ESPECIALIDAD) or "",
                    col(row, COL_RAZONSOCIAL) or col(row, COL_NOMBRE) or "",
                    col(row, COL_VALOROBRA) or "",
                    col(row, COL_PORCENTAJE) or "",
                    "",  # PORCENTAJE_CALCULADO
                    col(row, COL_SUBTOTAL) or "",
                    "",  # SUBTOTAL_CALCULADO
                    col(row, COL_TOTAL) or "",
                    "",  # TOTAL_CALCULADO
                    ", ".join(anomalies),
                    estado,
                ])
                counts[estado] += 1
                continue

            # ── Process row ─────────────────────────────────────────────────────
            try:
                payload, row_anomalies = self._build_payload(row, idx, orchestrator)
                anomalies = row_anomalies

                # ── Call cotizar (no insert) ──────────────────────────────────
                excel_total_raw = col(row, COL_TOTAL)
                excel_subtotal_raw = col(row, COL_SUBTOTAL)
                excel_pct_raw = col(row, COL_PORCENTAJE)

                excel_total = Decimal(str(excel_total_raw)) if excel_total_raw is not None else None
                excel_subtotal = Decimal(str(excel_subtotal_raw)) if excel_subtotal_raw is not None else None
                excel_pct = Decimal(str(excel_pct_raw)) if excel_pct_raw is not None else None

                # Build cotizacion input from the same payload
                cotizacion = orchestrator.cotizar_legacy_proceso(payload)

                pct_calculado = (cotizacion.porcentaje_liquidacion * Decimal(100)) if cotizacion else None
                calc_total = cotizacion.total if cotizacion else None
                calc_subtotal = cotizacion.total_subtotal if cotizacion else None

                # Exact comparisons — no tolerance
                if cotizacion is not None:
                    if excel_total is not None and cotizacion.total != excel_total:
                        anomalies.append("Total no coincide con cálculo")
                    if excel_subtotal is not None and cotizacion.total_subtotal != excel_subtotal:
                        anomalies.append("Subtotal no coincide con cálculo")
                    if excel_pct is not None and pct_calculado is not None and pct_calculado != excel_pct:
                        anomalies.append("Porcentaje no coincide con cálculo")

                # Determine ESTADO
                if anomalies:
                    estado = "AVISO"
                else:
                    estado = "OK"
                counts[estado] += 1

                # ── Expediente ────────────────────────────────────────────────
                nroexpdte = col(row, COL_NROEXPDTE)
                expediente = str(nroexpdte).strip() if nroexpdte else f"LEGACY-{idx}"
                if not expediente or expediente.upper() == "NULL":
                    expediente = f"LEGACY-{idx}"

                # ── Fecha ─────────────────────────────────────────────────────
                fecha_raw = col(row, COL_FECHA)
                fecha_str = ""
                if fecha_raw:
                    try:
                        fecha_str = _parse_fecha(fecha_raw).isoformat()
                    except ValueError:
                        anomalies.append(f"Fecha inválida: {fecha_raw}")

                # ── Razón social ───────────────────────────────────────────────
                razon_social = col(row, COL_RAZONSOCIAL)
                if not razon_social or str(razon_social).strip() == "":
                    razon_social = col(row, COL_NOMBRE)
                razon_social_str = str(razon_social).strip() if razon_social else ""

                writer.writerow([
                    fila,
                    nro_val if nro_val is not None else "",
                    expediente,
                    fecha_str,
                    (col(row, COL_ESPECIALIDAD) or "").strip(),
                    razon_social_str,
                    col(row, COL_VALOROBRA) or "",
                    excel_pct if excel_pct is not None else "",
                    str(pct_calculado) if pct_calculado is not None else "",
                    excel_subtotal if excel_subtotal is not None else "",
                    str(calc_subtotal) if calc_subtotal is not None else "",
                    excel_total if excel_total is not None else "",
                    str(calc_total) if calc_total is not None else "",
                    ", ".join(anomalies) if anomalies else "",
                    estado,
                ])

            except Exception as exc:  # noqa: BLE001
                estado = "ERROR"
                counts[estado] += 1
                self._log(self.style.ERROR(f"  ERROR row {idx}: {exc}"))
                writer.writerow([
                    fila,
                    nro_val if nro_val is not None else "",
                    f"LEGACY-{idx}",
                    "",
                    (col(row, COL_ESPECIALIDAD) or "").strip(),
                    (col(row, COL_RAZONSOCIAL) or col(row, COL_NOMBRE) or "").strip(),
                    col(row, COL_VALOROBRA) or "",
                    col(row, COL_PORCENTAJE) or "",
                    "",
                    col(row, COL_SUBTOTAL) or "",
                    "",
                    col(row, COL_TOTAL) or "",
                    "",
                    f"ERROR: {exc}",
                    "ERROR",
                ])

        out_file.close()

        # ── Summary ─────────────────────────────────────────────────────────────
        total = sum(counts.values())
        self._log(self.style.SUCCESS("\n=== REPORTE LEGACY DETALLADO ==="))
        self._log(f"  Total filas procesadas: {total}")
        for k, v in counts.items():
            self._log(f"  {k}: {v}")
        self._log(f"  CSV: {output_path}")

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _log(self, msg):
        """Write a message, handling Windows encoding quirks."""
        try:
            self.stdout.write(msg)
        except UnicodeEncodeError:
            safe = msg.encode("ascii", "replace").decode("ascii")
            self.stdout.write(safe)

    def _build_payload(self, row: list, idx: int, orchestrator) -> tuple:
        """
        Build a LiquidacionEdificacionesLegacyIn from an Excel row list.

        Returns (payload, anomalies_list).
        Mirrors the logic of import_legacy_edificaciones._build_payload.
        Raises ValueError on fatal parse errors.
        """
        anomalies: list[str] = []

        def col(n: int):
            return row[n] if len(row) > n else None

        # ── Null checks ─────────────────────────────────────────────────────────
        def _check_null(col_idx: int, col_name: str):
            val = row[col_idx]
            if val is None:
                anomalies.append(f"NULL en columna {col_name}")
                return True
            if str(val).strip().upper() == "NULL":
                anomalies.append(f"NULL en columna {col_name}")
                return True
            return False

        _check_null(COL_CODPAGO, "CODPAGO")
        _check_null(COL_DNI, "DNI")
        _check_null(COL_NOMBRE, "NOMBRE")
        _check_null(COL_DIRECCION, "DIRECCION")
        _check_null(COL_VALOROBRA, "VALOROBRA")
        _check_null(COL_FECHA, "FECHA")
        _check_null(COL_TOTAL, "TOTAL")
        _check_null(COL_SUBTOTAL, "SUBTOTAL")
        _check_null(COL_PORCENTAJE, "PORCENTAJE")
        _check_null(COL_NROREV, "NROREV")

        # ── Encoding issues ───────────────────────────────────────────────────
        for col_idx, col_name in [(COL_NOMBRE, "NOMBRE"), (COL_DIRECCION, "DIRECCION"), (COL_PROYECTO, "PROYECTO")]:
            val = row[col_idx]
            if val is not None and "�" in str(val):
                anomalies.append("Encoding issue en texto")

        # ── Especialidad ─────────────────────────────────────────────────────
        especialidad_raw = col(COL_ESPECIALIDAD)
        especialidad_str = str(especialidad_raw).strip() if especialidad_raw else ""
        valid_esp = {"TODAS", "Estructuras", "Inst. Sanitarias", "Inst. Mecánico Eléctricas"}
        if especialidad_str and especialidad_str not in valid_esp:
            anomalies.append(f"Especialidad no estándar: {especialidad_str}")

        # ── Entidad (documento) ───────────────────────────────────────────────
        dni_raw = col(COL_DNI)
        ruc_raw = col(COL_RUC)
        dni_str = str(dni_raw).strip() if dni_raw else ""
        ruc_str = str(ruc_raw).strip() if ruc_raw else ""

        if not dni_str and not ruc_str:
            anomalies.append("Sin documento (DNI y RUC faltantes)")

        tipo_documento, numero_documento = _resolve_tipo_documento(dni_raw)

        if not dni_str and ruc_str:
            tipo_documento = "RUC"
            numero_documento = ruc_str

        # ── municipalidad ─────────────────────────────────────────────────────
        codpago = col(COL_CODPAGO)
        codpago_str = str(codpago).strip() if codpago is not None else ""
        if not codpago_str:
            anomalies.append("CodPago faltante")
            municipalidad = Municipalidad.objects.first()
            if not municipalidad:
                raise ValueError("No hay municipalidades en la base de datos")
        else:
            try:
                municipalidad = Municipalidad.objects.get(codigo=codpago_str)
            except ObjectDoesNotExist:
                anomalies.append(f"CodPago no encontrado: {codpago_str}")
                municipalidad = Municipalidad.objects.first()
                if not municipalidad:
                    raise ValueError(f"No se encontró municipalidad para {codpago_str!r} y no hay fallback")

        municipalidad_id = municipalidad.id

        # ── distrito ─────────────────────────────────────────────────────────
        dptoprdri = col(COL_DPTOPRDI)
        distrito_nombre = str(dptoprdri).split("/")[-1].strip() if dptoprdri else ""
        distrito_nombre, provincia_nombre = _normalizar_distrito(distrito_nombre)
        distrito_qs = UbigeoDistrito.objects.filter(nombre__iexact=distrito_nombre)
        if provincia_nombre:
            distrito_qs = distrito_qs.filter(provincia__nombre__iexact=provincia_nombre)
        distrito = distrito_qs.first()
        if distrito:
            distrito_id = distrito.id
        else:
            fallback = UbigeoDistrito.objects.first()
            if fallback:
                distrito_id = fallback.id
                anomalies.append(f"Distrito no encontrado: {distrito_nombre}")
            else:
                raise ValueError("No UbigeoDistrito records found in database")

        # ── Razón social ─────────────────────────────────────────────────────
        razon_social = col(COL_RAZONSOCIAL)
        if razon_social is None or str(razon_social).strip() == "":
            razon_social = col(COL_NOMBRE)
        if razon_social is None or str(razon_social).strip() == "":
            raise ValueError("RAZONSOCIAL/NOMBRE is empty")
        razon_social_str = str(razon_social).strip()

        # ── Dirección ────────────────────────────────────────────────────────
        direccion = col(COL_DIRECCION)
        if direccion is None or str(direccion).strip() == "":
            raise ValueError("DIRECCION is empty")
        direccion_str = str(direccion).strip()

        # ── Expediente ───────────────────────────────────────────────────────
        nroexpdte = col(COL_NROEXPDTE)
        expediente = str(nroexpdte).strip() if nroexpdte else f"LEGACY-{idx}"
        if not expediente or expediente.upper() == "NULL":
            expediente = f"LEGACY-{idx}"

        # ── Fecha de registro ────────────────────────────────────────────────
        fecha_raw = col(COL_FECHA)
        try:
            fecha_registro = _parse_fecha(fecha_raw)
        except ValueError as exc:
            anomalies.append(f"Fecha inválida: {fecha_raw}")
            raise ValueError(f"Invalid FECHA {fecha_raw!r}: {exc}") from exc

        # ── Valor de obra ────────────────────────────────────────────────────
        valor_raw = col(COL_VALOROBRA)
        if valor_raw is None or str(valor_raw).strip() == "":
            raise ValueError("VALOROBRA is empty")
        try:
            valor_declarado = Decimal(str(valor_raw).strip())
        except (InvalidOperation, ValueError) as exc:
            anomalies.append(f"ValorObra inválido o cero: {valor_raw}")
            raise ValueError(f"Invalid VALOROBRA {valor_raw!r}: {exc}") from exc

        if valor_declarado <= 0:
            anomalies.append(f"ValorObra inválido o cero: {valor_raw}")
            raise ValueError(f"VALOROBRA must be positive, got {valor_declarado}")

        # ── Número de revisión ────────────────────────────────────────────────
        nrorev_raw = col(COL_NROREV)
        nrorev_val: int | None = None
        if nrorev_raw is not None and str(nrorev_raw).strip() != "":
            try:
                nrorev_val = int(float(str(nrorev_raw).strip()))
            except (ValueError, TypeError):
                nrorev_val = None

        if nrorev_val is None:
            anomalies.append("NROREV faltante — se asigna 1")
            nrorev_val = 1
        elif nrorev_val > MAX_REVISIONES:
            anomalies.append(f"NROREV anómalo: {nrorev_val}")

        # ── Denominación de proyecto ─────────────────────────────────────────
        proyecto_raw = col(COL_PROYECTO)
        denominacion_de_proyecto = str(proyecto_raw).strip() if proyecto_raw else None

        # ── Número secuencial ─────────────────────────────────────────────────
        nro_raw = col(COL_NRO)
        numero: int | None = None
        if nro_raw is not None and str(nro_raw).strip() != "":
            try:
                numero = int(float(str(nro_raw).strip()))
            except (ValueError, TypeError):
                numero = None

        # ── Build schema objects ──────────────────────────────────────────────
        proyecto_schema = ProyectoCotizarSchema(
            denominacion=razon_social_str,
            nombre_propietario=razon_social_str,
            direccion=direccion_str,
            distrito_id=distrito_id,
            entidad=EntidadInlineSchema(
                tipo_documento=tipo_documento,
                numero_documento=numero_documento,
                razon_social=razon_social_str,
            ),
        )

        liquidacion_general_partial = LiquidacionGeneralLegacyIn(
            municipalidad_id=municipalidad_id,
            expediente=expediente,
            observacion=None,
            retencion=False,
            proyecto=proyecto_schema,
            contacto=None,
            fecha_registro=fecha_registro,
            denominacion_de_proyecto_liquidacion=denominacion_de_proyecto,
            descripcion_legacy=None,
        )

        liquidacion_especifica = LiquidacionPorcentajeObraIn(
            datos=LiquidacionPorcentajeObraDatosIn(valor_declarado=valor_declarado),
            tarifas=[],  # auto-fill mode
        )

        payload = LiquidacionEdificacionesLegacyIn(
            liquidacion_general=liquidacion_general_partial,
            liquidacion_especifica=liquidacion_especifica,
            numero_revision=nrorev_val,
            numero=numero,
        )

        return payload, anomalies
